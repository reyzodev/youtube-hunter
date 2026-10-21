"""app.core.exports - Youtube Scrapper."""

import csv
import json
import os
import re
import sqlite3

# Control characters that Excel/openpyxl reject outright (everything except
# \t \n \r). YouTube titles/descriptions/transcripts DO occasionally contain
# them — previously this made the export blow up at the very end of a scrape.
_ILLEGAL_CELL_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def _clean_cell(value):
    if isinstance(value, str):
        return _ILLEGAL_CELL_RE.sub("", value)
    return value


def write_export(data, path, fmt):
    """Write scraped rows to disk in the chosen format.

    Hardened: values are sanitized (illegal control chars stripped) and the
    file is written atomically (temp file + os.replace) so a crash or 'disk
    full' mid-write can never leave a corrupt half-file behind."""
    if not data:
        return
    data = [{k: _clean_cell(v) for k, v in row.items()} for row in data]
    tmp = path + ".tmp"
    try:
        if fmt == "csv":
            with open(tmp, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.DictWriter(f, fieldnames=data[0].keys())
                writer.writeheader()
                writer.writerows(data)
        elif fmt in ("json", "jsonl"):   # Niche / Archive pass "jsonl"
            with open(tmp, "w", encoding="utf-8") as f:
                for row in data:
                    f.write(json.dumps(row, ensure_ascii=False) + "\n")
        elif fmt == "xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "videos"
            headers = list(data[0].keys())
            ws.append(headers)
            for row in data:
                ws.append([row.get(h, "") for h in headers])
            wb.save(tmp)
        elif fmt == "sqlite":
            conn = sqlite3.connect(tmp)
            cur = conn.cursor()
            headers = list(data[0].keys())
            cols = ", ".join([f'"{h}" TEXT' for h in headers])
            cur.execute("DROP TABLE IF EXISTS videos")
            cur.execute(f"CREATE TABLE videos ({cols})")
            placeholders = ", ".join(["?"] * len(headers))
            for row in data:
                cur.execute(f"INSERT INTO videos VALUES ({placeholders})",
                            [str(row.get(h, "")) for h in headers])
            conn.commit()
            conn.close()
        os.replace(tmp, path)
    finally:
        try:
            if os.path.exists(tmp):
                os.remove(tmp)
        except Exception:
            pass
def read_checkpoint_rows(path):
    """Read rows back from a JSONL checkpoint file. Returns a list of dicts
    (empty on any problem — resume is always best-effort)."""
    rows = []
    try:
        with open(path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    row = json.loads(line)
                    if isinstance(row, dict) and row.get("video_id"):
                        rows.append(row)
                except Exception:
                    continue
    except Exception:
        return []
    return rows
